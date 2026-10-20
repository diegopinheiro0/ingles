"""Confere a planilha do guia de pronúncia antes de gerar o PDF.

Uso: python validar.py Guia_Ingles_Pronuncia.xlsx
Saída: lista de ERROS (impedem a entrega) e AVISOS (revisar). Exit 1 se houver erro.
"""
import re
import sys
from collections import Counter

from openpyxl import load_workbook

ABAS = ["Textos", "Legenda", "Seções", "Grupos", "Palavras", "Pares mínimos", "Contrastes", "Falsos cognatos", "Teste"]
PRIORIDADES = {"Alta", "Média", "Pontual"}
STATUS = {"", "treinar", "dominada"}
# Fora da notação: vogais nasais, í/ú, IPA (exceto ɾ e o t' que é ASCII)
PROIBIDOS_PRON = set("ãõĩũÃÕĨŨíúÍÚìùàèòçÇüÜəɪʊʃʒŋθðɑɔʌɚɝɹʔ")
LETRAS_PT_IPA = set("áéíóúâêôãõçÁÉÍÓÚÂÊÔ")

erros, avisos = [], []

def linhas(ws):
    it = ws.iter_rows(values_only=True)
    cab = [str(h).strip() if h else "" for h in next(it)]
    out = []
    for n, r in enumerate(it, start=2):
        if all(v in (None, "") for v in r):
            continue
        d = {h: ("" if v is None else str(v).strip()) for h, v in zip(cab, r)}
        d["_linha"] = n
        out.append(d)
    return out

def pron_principal(p):
    return p.split(" (")[0]

def checa_pron(onde, pron):
    ruins = sorted({c for c in pron if c in PROIBIDOS_PRON})
    if ruins:
        erros.append(f"{onde}: pronúncia com caracteres fora da notação: {' '.join(ruins)} → '{pron}'")
    principal = pron_principal(pron)
    if re.search(r"(^|[-\s])U[ÓÁÉÊ]", principal):
        avisos.append(f"{onde}: 'U' + vogal parece W escrito como U (use W): '{pron}'")
    if re.search(r"[óÓ](?![-\s]?[rR])(?![iI])", principal):
        erros.append(f"{onde}: 'ó' sem R depois; o guia usa 'á' (law LÁ, thought THÁT). 'ó' só em ór e ói: '{pron}'")
    if re.search(r"(^|[-\s])IU", principal, re.I):
        avisos.append(f"{onde}: 'IU' parece Y escrito como I (use YUU/yuu): '{pron}'")
    silabas = [s for s in re.split(r"[-\s]", principal) if s]
    tonicas = [s for s in silabas if s.upper() == s and any(ch.isalpha() for ch in s)]
    if len(silabas) > 1 and not tonicas:
        avisos.append(f"{onde}: nenhuma sílaba em MAIÚSCULAS (tônica): '{pron}'")

def checa_ipa(onde, ipa):
    if not ipa:
        avisos.append(f"{onde}: IPA vazio")
        return
    if "/" in ipa:
        erros.append(f"{onde}: IPA sem barras (o PDF acrescenta): '{ipa}'")
    if any(c in LETRAS_PT_IPA for c in ipa):
        erros.append(f"{onde}: IPA com letra acentuada do português: '{ipa}'")

def main(path):
    wb = load_workbook(path, data_only=True)
    faltando = [a for a in ABAS if a not in wb.sheetnames]
    if faltando:
        erros.append(f"Abas ausentes: {', '.join(faltando)}")
        return
    secoes = {s["Seção"] for s in linhas(wb["Seções"])}
    grupos = linhas(wb["Grupos"])
    chaves = Counter(g["Chave"] for g in grupos)
    for k, n in chaves.items():
        if n > 1:
            erros.append(f"Grupos: chave '{k}' repetida {n} vezes")
    for g in grupos:
        onde = f"Grupos linha {g['_linha']} ({g['Chave']})"
        for campo in ("Chave", "Seção", "Nome", "Som", "Erro comum", "Frase de treino", "Prioridade"):
            if not g.get(campo):
                erros.append(f"{onde}: campo '{campo}' vazio")
        if g["Seção"] and g["Seção"] not in secoes:
            erros.append(f"{onde}: seção '{g['Seção']}' não existe na aba Seções")
        if g["Prioridade"] and g["Prioridade"] not in PRIORIDADES:
            erros.append(f"{onde}: prioridade '{g['Prioridade']}' (use Alta, Média ou Pontual)")
        if not g.get("Fonte"):
            avisos.append(f"{onde}: sem fonte")

    palavras = linhas(wb["Palavras"])
    por_grupo = {}
    for w in palavras:
        onde = f"Palavras linha {w['_linha']} ({w['Palavra'] or '?'})"
        if w["Grupo"] not in chaves:
            erros.append(f"{onde}: grupo '{w['Grupo']}' não existe na aba Grupos")
        if not w["Palavra"] or not w["Pronúncia"]:
            erros.append(f"{onde}: palavra ou pronúncia vazia")
            continue
        checa_pron(onde, w["Pronúncia"])
        checa_ipa(onde, w["IPA (EUA)"])
        if w["Status"] not in STATUS:
            erros.append(f"{onde}: status '{w['Status']}' (use treinar ou dominada)")
        if not w["Fonte"]:
            avisos.append(f"{onde}: sem fonte")
        por_grupo.setdefault(w["Grupo"], []).append(w["Palavra"].lower())
    dup = Counter((w["Grupo"], w["Palavra"].lower()) for w in palavras)
    for (g, p), n in dup.items():
        if n > 1:
            erros.append(f"Palavras: '{p}' repetida {n} vezes no grupo {g}")
    pron_por_palavra = {}
    for w in palavras:
        pron_por_palavra.setdefault(w["Palavra"].lower(), set()).add(pron_principal(w["Pronúncia"]))
    for p, prons in pron_por_palavra.items():
        if len(prons) > 1:
            avisos.append(f"Palavras: '{p}' aparece com pronúncias diferentes: {' | '.join(sorted(prons))}")
    for g in grupos:
        if g["Chave"] not in por_grupo:
            avisos.append(f"Grupos: '{g['Chave']}' sem nenhuma palavra")

    contrastes = {c["Contraste"] for c in linhas(wb["Contrastes"])}
    for r in linhas(wb["Pares mínimos"]):
        onde = f"Pares mínimos linha {r['_linha']}"
        if r["Contraste"] not in contrastes:
            avisos.append(f"{onde}: contraste '{r['Contraste']}' sem dica na aba Contrastes")
        for c in "ABC":
            if r.get(f"Palavra {c}") and not r.get(f"Pronúncia {c}"):
                erros.append(f"{onde}: palavra {c} sem pronúncia")
            if r.get(f"Pronúncia {c}"):
                checa_pron(onde, r[f"Pronúncia {c}"])

    falsos = linhas(wb["Falsos cognatos"])
    for k, n in Counter(f["Palavra"].lower() for f in falsos).items():
        if n > 1:
            erros.append(f"Falsos cognatos: '{k}' repetida {n} vezes")
    for f in falsos:
        onde = f"Falsos cognatos linha {f['_linha']} ({f['Palavra']})"
        for campo in ("Tabela", "Palavra", "Pronúncia", "Parece", "Significa", "Para dizer o que parece", "Frase"):
            if not f.get(campo):
                erros.append(f"{onde}: campo '{campo}' vazio")
        if f["Pronúncia"]:
            checa_pron(onde, f["Pronúncia"])
        checa_ipa(onde, f.get("IPA (EUA)", ""))

    for t in linhas(wb["Teste"]):
        if t["Gabarito"] not in ("a", "b"):
            erros.append(f"Teste linha {t['_linha']}: gabarito '{t['Gabarito']}' (use a ou b)")

    print(f"{len(grupos)} grupos, {len(palavras)} palavras, {len(falsos)} falsos cognatos")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1])
    for e in erros:
        print("ERRO  ", e)
    for a in avisos:
        print("AVISO ", a)
    print(f"Resultado: {len(erros)} erro(s), {len(avisos)} aviso(s)")
    sys.exit(1 if erros else 0)
