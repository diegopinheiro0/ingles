"""Consulta a pronúncia americana (IPA US) no Cambridge Dictionary e confere a planilha.

Uso:
  python cambridge.py palavra [palavra ...]           → mostra o IPA US de cada palavra
  python cambridge.py --conferir Guia_Ingles_Pronuncia.xlsx
                                                     → compara o IPA da planilha com o Cambridge

Consultas ficam em cache em .cache_cambridge.json (não vai para o repositório).
Na comparação, conta a vogal tônica e a posição da tônica; diferenças de ː, de pontos de sílaba
e de t̬ (= ɾ no guia) são ignoradas.
"""
import html
import json
import os
import re
import sys
import time
import urllib.parse
import urllib.request

BASE = "https://dictionary.cambridge.org/dictionary/english/"
CACHE = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".cache_cambridge.json")
UA = {"User-Agent": "Mozilla/5.0 (guia-pronuncia)"}

def _cache():
    try:
        return json.load(open(CACHE, encoding="utf-8"))
    except (OSError, ValueError):
        return {}

def _limpa_ipa(frag):
    frag = re.sub(r'<span class="sp dsp">(.*?)</span>', r"(\1)", frag)  # sons opcionais
    return html.unescape(re.sub(r"<[^>]+>", "", frag)).strip()

def buscar(palavra, cache=None):
    """Lista de IPAs US (sem barras) da página da palavra; [] se não houver."""
    cache = _cache() if cache is None else cache
    chave = palavra.lower().strip()
    if chave in cache:
        return cache[chave]
    url = BASE + urllib.parse.quote(chave.replace(" ", "-"))
    req = urllib.request.Request(url, headers=UA)
    try:
        with urllib.request.urlopen(req, timeout=25) as r:
            final = r.geturl()
            pagina = r.read().decode("utf-8", "replace")
    except Exception as e:  # rede bloqueada, 404 etc.
        print(f"  ! {palavra}: {e}", file=sys.stderr)
        return []
    ipas = []
    if final.rstrip("/").endswith("/english"):  # palavra não encontrada: redireciona para a busca
        ipas = []
    else:
        for m in re.finditer(r'<span class="us dpron-i[^>]*>(.*?)</span>\s*/\s*</span>', pagina, re.S):
            ipa = re.search(r'<span class="ipa[^"]*">(.*)', m.group(1), re.S)
            if ipa:
                v = _limpa_ipa(ipa.group(1))
                for parte in v.split(","):
                    parte = parte.strip().strip("-")
                    if parte and parte not in ipas:
                        ipas.append(parte)
    cache[chave] = ipas
    json.dump(cache, open(CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=0)
    time.sleep(0.4)
    return ipas

# ---------------------------------------------------------------- comparação
VOGAIS = [("eɪ", "eɪ"), ("oʊ", "oʊ"), ("aɪ", "aɪ"), ("aʊ", "aʊ"), ("ɔɪ", "ɔɪ"), ("ɝ", "ɝ"), ("ɚ", "ɚ"),
          ("ɪ", "ɪ"), ("i", "i"), ("ɛ", "ɛ"), ("e", "ɛ"), ("æ", "æ"), ("ɑ", "ɑ"), ("ɔ", "ɔ"), ("ʌ", "ʌ"),
          ("ə", "ə"), ("ʊ", "ʊ"), ("u", "u")]

def normaliza(ipa):
    ipa = re.sub(r"\(.*?\)", "", ipa)  # sons opcionais fora
    return ipa.replace("ː", "").replace(".", "").replace("·", "").replace("t̬", "ɾ").replace("ˌ", "").replace(" ", "")

def vogais(ipa):
    s, out, i, forte = normaliza(ipa), [], 0, False
    while i < len(s):
        if s[i] == "ˈ":
            forte = True; i += 1; continue
        for k, v in VOGAIS:
            if s.startswith(k, i):
                if v == "ɔ" and s[i + 1:i + 2] != "r" and k != "ɔɪ":
                    v = "ɑ"  # guia (opção B): ɔ só antes de R; o resto é ɑ
                if v == "ɝ":
                    v = "ɚ"
                out.append((v, forte)); forte = False; i += len(k); break
        else:
            i += 1
    if len(out) == 1:
        out[0] = (out[0][0], True)
    return out

def tonica(vs):
    return next((v for v, f in vs if f), None)

def compativel(guia, camb):
    g, c = vogais(guia), vogais(camb)
    if tonica(g) != tonica(c):
        return False, "vogal tônica"
    if len(g) == len(c) and [f for _, f in g] != [f for _, f in c]:
        return False, "posição da tônica"
    return True, ""

def conferir(xlsx):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, data_only=True)
    itens = [("Palavras", r[0], r[2], r[3], r[4]) for r in wb["Palavras"].iter_rows(min_row=2, values_only=True) if r[2]]
    itens += [("Falsos", "", r[1], r[2], r[3]) for r in wb["Falsos cognatos"].iter_rows(min_row=2, values_only=True) if r[1]]
    cache, div, sem = _cache(), [], []
    for aba, grupo, palavra, pron, ipa in itens:
        chave = re.sub(r"\(.*?\)", "", palavra).strip()
        camb = buscar(chave, cache)
        if not camb:
            sem.append(palavra); continue
        if not ipa:
            div.append(f"{palavra} ({grupo or aba}): sem IPA na planilha; Cambridge US: {' | '.join(camb)}"); continue
        motivos = [compativel(ipa, c) for c in camb]
        if not any(ok for ok, _ in motivos):
            div.append(f"{palavra} ({grupo or aba}): {pron} /{ipa}/ × Cambridge US /{' | '.join(camb)}/ [{motivos[0][1]}]")
    print(f"{len(itens)} entradas conferidas, {len(div)} divergência(s), {len(sem)} sem página no Cambridge")
    for d in div:
        print("  DIVERGE ", d)
    if sem:
        print("  SEM PÁGINA:", ", ".join(sem))
    return div

if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    if sys.argv[1] == "--conferir":
        conferir(sys.argv[2])
    else:
        for p in sys.argv[1:]:
            print(p, "→", " | ".join(f"/{i}/" for i in buscar(p)) or "não encontrada")
