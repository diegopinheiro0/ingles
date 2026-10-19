"""Gera o PDF do guia a partir da planilha mestre.

Uso: python gerar_pdf.py Guia_Ingles_Pronuncia.xlsx [saida.pdf]
Requer: pip install openpyxl reportlab
"""
import os
import sys
from collections import OrderedDict
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

# ---------------------------------------------------------------- fontes (precisam cobrir IPA: ɪ ʊ ə æ ɾ ʃ θ ð ŋ)
FONT_DIRS = ["/usr/share/fonts/truetype/dejavu", "C:/Windows/Fonts", "/Library/Fonts",
             os.path.expanduser("~/Library/Fonts"), os.path.dirname(os.path.abspath(__file__))]
def find_font(names):
    for d in FONT_DIRS:
        for n in names:
            p = os.path.join(d, n)
            if os.path.exists(p):
                return p
    sys.exit(f"Fonte não encontrada: {names}. Instale DejaVu Sans ou coloque o .ttf na pasta do script.")

pdfmetrics.registerFont(TTFont("Body", find_font(["DejaVuSans.ttf", "segoeui.ttf", "arial.ttf"])))
pdfmetrics.registerFont(TTFont("Bold", find_font(["DejaVuSans-Bold.ttf", "segoeuib.ttf", "arialbd.ttf"])))

NAVY = colors.HexColor("#1F4E79")
LIGHT = colors.HexColor("#EAF1F8")
GRID = colors.HexColor("#B8C7D9")
GREY = colors.HexColor("#6B7785")
AMBER = colors.HexColor("#FFF4D6")

W = A4[0] - 30 * mm

def style(name, size=9, font="Body", color=colors.black, leading=None, **kw):
    return ParagraphStyle(name, fontName=font, fontSize=size, leading=leading or size * 1.3,
                          textColor=color, alignment=TA_LEFT, **kw)

ST = {
    "title": style("title", 20, "Bold", NAVY, 25),
    "sub": style("sub", 10, color=GREY),
    "h1": style("h1", 15, "Bold", NAVY, 19, spaceBefore=4, spaceAfter=3),
    "h2": style("h2", 11, "Bold", NAVY, 14, spaceBefore=6, spaceAfter=2),
    "intro": style("intro", 9, color=colors.HexColor("#333333"), spaceAfter=6),
    "cell": style("cell", 8.5),
    "cellb": style("cellb", 8.5, "Bold"),
    "ipa": style("ipa", 8, color=GREY),
    "head": style("head", 10, "Bold", colors.white),
    "headsm": style("headsm", 8.5, "Bold", colors.white),
    "sub_label": style("sub_label", 8, "Bold", NAVY),
    "label": style("label", 8, "Bold", NAVY),
    "frase": style("frase", 8.5, color=colors.HexColor("#14365A")),
    "small": style("small", 7.5, color=GREY),
}

def p(text, st="cell"):
    return Paragraph(escape(str(text or "")).replace("\n", "<br/>"), ST[st])

def pm(markup, st="cell"):
    return Paragraph(markup, ST[st])

# ---------------------------------------------------------------- leitura da planilha
def rows(ws):
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h else "" for h in next(it)]
    out = []
    for r in it:
        if all(v in (None, "") for v in r):
            continue
        out.append({h: ("" if v is None else str(v).strip()) for h, v in zip(header, r)})
    return out

def load(path):
    wb = load_workbook(path, data_only=True)
    d = {name: rows(wb[name]) for name in wb.sheetnames}
    d["textos"] = {r["Chave"]: r["Texto"] for r in d["Textos"]}
    return d

# ---------------------------------------------------------------- blocos
def grid_style(extra=()):
    return TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, GRID),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        *extra,
    ])

def simple_table(header, data, widths, header_style="headsm"):
    t = Table([[p(h, header_style) for h in header]] + data, colWidths=widths, repeatRows=1)
    t.setStyle(grid_style([("BACKGROUND", (0, 0), (-1, 0), NAVY)]))
    return t

def word_cells(w):
    if not w:
        return ["", "", ""]
    return [p(w["Palavra"], "cellb"), p(w["Pronúncia"]), p(f"/{w['IPA (EUA)']}/" if w["IPA (EUA)"] else "", "ipa")]

def card(g, words):
    cw = [27 * mm, 34 * mm, 29 * mm] * 2
    data = [[pm(f"{escape(g['Nome'])} = {escape(g['Som'])}", "head")] + [""] * 5]
    cmds = [("SPAN", (0, 0), (-1, 0)), ("BACKGROUND", (0, 0), (-1, 0), NAVY)]
    subs = OrderedDict()
    for w in words:
        subs.setdefault(w["Subgrupo"], []).append(w)
    for sub, ws in subs.items():
        if sub:
            r = len(data)
            data.append([p(sub, "sub_label")] + [""] * 5)
            cmds += [("SPAN", (0, r), (-1, r)), ("BACKGROUND", (0, r), (-1, r), LIGHT)]
        for i in range(0, len(ws), 2):
            pair = ws[i:i + 2] + [None]
            data.append(word_cells(pair[0]) + word_cells(pair[1]))
    for label, key, st in (("Erro comum", "Erro comum", "cell"), ("Atenção", "Atenção", "cell"),
                           ("Frase", "Frase de treino", "frase")):
        if g.get(key):
            r = len(data)
            data.append([p(label, "label"), p(g[key], st)] + [""] * 4)
            cmds += [("SPAN", (1, r), (-1, r)), ("BACKGROUND", (0, r), (0, r), LIGHT)]
    t = Table(data, colWidths=cw)
    t.setStyle(grid_style(cmds))
    return KeepTogether([t, Spacer(1, 7)])

# ---------------------------------------------------------------- documento
def build(xlsx, out):
    d = load(xlsx)
    tx = d["textos"]
    grupos = d["Grupos"]
    palavras = d["Palavras"]
    by_group = OrderedDict((g["Chave"], []) for g in grupos)
    for w in palavras:
        by_group.setdefault(w["Grupo"], []).append(w)
    sections = sorted(d["Seções"], key=lambda s: float(s["Ordem"] or 0))

    story = [p(tx.get("titulo", ""), "title"), Spacer(1, 3), p(tx.get("subtitulo", ""), "sub"), Spacer(1, 8)]

    # Legenda
    story += [p("Como ler a pronúncia", "h1"), p(tx.get("legenda_intro", ""), "intro")]
    leg = [[p(r["Símbolo"], "cellb"), p(r["IPA"], "ipa"), p(r["Como fazer"]), p(r["Exemplo"]), p(r["Erro comum"], "small")]
           for r in d["Legenda"]]
    story.append(simple_table(["Símbolo", "IPA", "Como fazer", "Exemplo", "Erro comum"], leg,
                              [28 * mm, 16 * mm, 52 * mm, 42 * mm, 42 * mm]))

    # Prioridades
    story += [Spacer(1, 8), p("Por onde começar", "h1")]
    pri = []
    for nivel in ("Alta", "Média", "Pontual"):
        nomes = [g["Nome"] for g in grupos if g["Prioridade"] == nivel]
        if nomes:
            pri.append([p(nivel, "cellb"), p(" · ".join(nomes)), p(tx.get(f"prioridade_{nivel}", ""), "small")])
    story.append(simple_table(["Prioridade", "Grupos", "Por quê"], pri, [22 * mm, 108 * mm, 50 * mm]))

    # Método
    story += [Spacer(1, 8), p("Método de treino (10 minutos por dia)", "h1")]
    for i, line in enumerate([l for l in tx.get("metodo", "").split("\n") if l.strip()], 1):
        story.append(p(f"{i}. {line.strip()}", "cell"))
    story += [Spacer(1, 3), p(tx.get("ritmo", ""), "intro")]

    # Seções de grupos
    for n, s in enumerate(sections, 1):
        gs = [g for g in grupos if g["Seção"] == s["Seção"]]
        if not gs:
            continue
        story += [PageBreak(), p(f"{n}. {s['Seção']}", "h1"), p(s["Introdução"], "intro")]
        for g in gs:
            story.append(card(g, by_group.get(g["Chave"], [])))
    orphans = [g for g in grupos if g["Seção"] not in {s["Seção"] for s in sections}]
    if orphans:
        story += [PageBreak(), p("Outros grupos", "h1")] + [card(g, by_group.get(g["Chave"], [])) for g in orphans]

    # Pares mínimos
    story += [PageBreak(), p("Pares mínimos", "h1"), p(tx.get("pares_intro", ""), "intro")]
    dicas = {r["Contraste"]: r["Dica"] for r in d.get("Contrastes", [])}
    pares = OrderedDict()
    for r in d["Pares mínimos"]:
        pares.setdefault(r["Contraste"], []).append(r)
    for contraste, rs in pares.items():
        has_c = any(r["Palavra C"] for r in rs)
        cols = ["A", "B", "C"] if has_c else ["A", "B"]
        data = [[pm(f"<font name='Bold'>{escape(r['Palavra ' + c])}</font>  {escape(r['Pronúncia ' + c])}")
                 if r["Palavra " + c] else "" for c in cols] for r in rs]
        t = Table(data, colWidths=[W / len(cols)] * len(cols))
        t.setStyle(grid_style())
        story.append(KeepTogether([p(contraste, "h2"), p(dicas.get(contraste, ""), "intro"), t, Spacer(1, 4)]))

    # Falsos cognatos
    story += [PageBreak(), p("Falsos cognatos", "h1"), p(tx.get("falsos_intro", ""), "intro")]
    tabelas = OrderedDict()
    for r in d["Falsos cognatos"]:
        tabelas.setdefault(r["Tabela"], []).append(r)
    for nome, rs in tabelas.items():
        data = [[pm(f"<font name='Bold'>{escape(r['Palavra'])}</font><br/>{escape(r['Pronúncia'])}"
                    + (f"<br/><font color='#6B7785' size='7.5'>/{escape(r['IPA (EUA)'])}/</font>" if r["IPA (EUA)"] else "")),
                 p(r["Parece"]), p(r["Significa"]), p(r["Para dizer o que parece"]), p(r["Frase"], "frase")] for r in rs]
        story += [p(nome, "h2"),
                  simple_table(["Palavra", "Parece", "Significa", "Para dizer o que parece", "Frase"], data,
                               [32 * mm, 26 * mm, 36 * mm, 36 * mm, 50 * mm]), Spacer(1, 4)]
    if tx.get("armadilhas"):
        box = Table([[p(tx["armadilhas"])]], colWidths=[W])
        box.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), AMBER), ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#E0B84C")),
                                 ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                                 ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        story += [Spacer(1, 4), box]

    # Mini teste
    teste = d.get("Teste", [])
    if teste:
        story += [PageBreak(), p("Mini teste: falsos cognatos", "h1"), p(tx.get("teste_intro", ""), "intro")]
        for i, r in enumerate(teste, 1):
            story.append(KeepTogether([p(f"{i}. {r['Frase em português']}", "cellb"),
                                       p(f"     a) {r['Opção a']}"), p(f"     b) {r['Opção b']}"), Spacer(1, 4)]))
        gab = " · ".join(f"{i}{r['Gabarito']}" for i, r in enumerate(teste, 1))
        story += [Spacer(1, 10), pm(f"<font name='Bold'>Gabarito:</font> {escape(gab)}"), Spacer(1, 3),
                  p(tx.get("teste_rodape", ""), "small")]

    # Cola rápida
    story += [PageBreak(), p("Cola rápida", "h1"), p(tx.get("cola_intro", ""), "intro")]
    cola = []
    ordem = {sec["Seção"]: i for i, sec in enumerate(sections)}
    for g in sorted(grupos, key=lambda g: ordem.get(g["Seção"], len(ordem))):
        ex = ", ".join(f"{w['Palavra']} ({w['Pronúncia']})" for w in by_group.get(g["Chave"], [])[:2])
        cola.append([p(g["Nome"], "cellb"), p(g["Som"]), p(ex)])
    story.append(simple_table(["Grupo", "Som", "Exemplos"], cola, [40 * mm, 52 * mm, 88 * mm]))

    titulo = tx.get("titulo", "")
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Body", 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(15 * mm, 9 * mm, titulo)
        canvas.drawRightString(A4[0] - 15 * mm, 9 * mm, str(doc.page))
        canvas.restoreState()

    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=14 * mm, bottomMargin=16 * mm, title=titulo, author="Diego")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    print(f"PDF gerado: {out} ({len(grupos)} grupos, {len(palavras)} palavras)")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = sys.argv[1]
    build(src, sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(src)[0] + ".pdf")
