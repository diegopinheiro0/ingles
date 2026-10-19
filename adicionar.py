"""Acrescenta linhas aprovadas à planilha do guia, sem mexer no que já existe.

Uso: python adicionar.py Guia_Ingles_Pronuncia.xlsx novos.json [saida.xlsx]

novos.json — cada chave é o nome de uma aba; cada item usa os nomes das colunas do cabeçalho:
{
  "Grupos":   [{"Chave": "EA", "Seção": "Vogais com grafia traiçoeira", "Nome": "EA", "Som": "...", ...}],
  "Palavras": [{"Grupo": "EA", "Subgrupo": "", "Palavra": "bread", "Pronúncia": "BRÉD",
                "IPA (EUA)": "brɛd", "Fonte": "Vídeo X (out/2026)", "Status": "treinar"}],
  "Falsos cognatos": [...], "Pares mínimos": [...], "Contrastes": [...], "Teste": [...]
}
Coluna ausente no item fica vazia; Status vazio vira "treinar".
"""
import json
import sys
from copy import copy

from openpyxl import load_workbook

def main(xlsx, novos_json, out):
    wb = load_workbook(xlsx)
    novos = json.load(open(novos_json, encoding="utf-8"))
    for aba, itens in novos.items():
        if aba not in wb.sheetnames:
            sys.exit(f"Aba inexistente: {aba}")
        ws = wb[aba]
        cab = [c.value for c in ws[1]]
        desconhecidas = {k for it in itens for k in it} - set(cab)
        if desconhecidas:
            sys.exit(f"{aba}: colunas desconhecidas {sorted(desconhecidas)}. Cabeçalho: {cab}")
        ultima = ws.max_row
        while ultima > 1 and all(c.value in (None, "") for c in ws[ultima]):
            ultima -= 1
        modelo = ws[ultima] if ultima > 1 else None
        for i, it in enumerate(itens, start=1):
            if "Status" in cab and not it.get("Status"):
                it = {**it, "Status": "treinar"}
            for j, col in enumerate(cab, start=1):
                cel = ws.cell(row=ultima + i, column=j, value=it.get(col, "") or None)
                if modelo is not None:
                    cel.alignment = copy(modelo[j - 1].alignment)
                    cel.font = copy(modelo[j - 1].font)
        print(f"{aba}: +{len(itens)} linha(s)")
    wb.save(out)
    print(f"Salvo: {out}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else sys.argv[1])
